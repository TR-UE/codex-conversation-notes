"""Source-positioned Markdown to native Word structure and matching paginated PDF."""
from __future__ import annotations
import base64
from io import BytesIO
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageOps

ROOT=Path(__file__).resolve().parent
PALETTE={'yellow':'F6E7A9','rose':'F5D0D8','green':'CCE5D3','blue':'CEDFF5','purple':'DFD2EF'}
ROLES={'user':'你','assistant':'Codex','excerpt':'原文摘录'}


def clean(value): return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))
def slice16(value,a,b): return value.encode('utf-16-le')[a*2:b*2].decode('utf-16-le')
def span(node): return node['position']['start']['offset'],node['position']['end']['offset']
def element(tag,**attrs):
    node=OxmlElement('w:'+tag)
    for k,v in attrs.items():node.set(qn('w:'+k),str(v))
    return node
def shading(properties,color):
    old=properties.find(qn('w:shd'))
    if old is not None: properties.remove(old)
    properties.append(element('shd',val='clear',fill=color))


def prepare(documents):
    config=json.loads((ROOT/'export-runtime.json').read_text())
    node=config.get('node')
    if not node or not Path(node).is_file():raise ValueError('富格式导出环境未配置，请更新墨记')
    process=subprocess.run([node,str(ROOT/'export/export-engine.mjs'),str(ROOT/'export-runtime.json')],
        input=json.dumps({'documents':documents},ensure_ascii=False),capture_output=True,text=True,timeout=90)
    if process.returncode:raise ValueError('公式或图片准备失败：'+process.stderr.strip()[-350:])
    return json.loads(process.stdout)


class WordWriter:
    def __init__(self,appearance=None):
        self.doc=Document();self.appearance=appearance or {};self.bookmark=0
        section=self.doc.sections[0]
        section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
        section.top_margin=section.bottom_margin=Inches(.7)
        section.left_margin=section.right_margin=Inches(.78)
        self.width=(section.page_width-section.left_margin-section.right_margin)/12700
        self.header_fill={'paper':'F1EDE6','mist':'EAF0F8'}.get(self.appearance.get('style'),'F1F2F4')
        for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Heading 5','Heading 6','Comment Text','Caption']:
            if name not in self.doc.styles:continue
            style=self.doc.styles[name];style.font.name='Heiti SC';style.font.color.rgb=RGBColor.from_string('202020')
            fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
            for attribute in list(fonts.attrib):
                if attribute.rsplit('}',1)[-1].lower().endswith('theme'):del fonts.attrib[attribute]
            fonts.set(qn('w:eastAsia'),'Heiti SC')
            if name.startswith(('Title','Heading','Subtitle')):style.font.color.rgb=RGBColor(0,0,0)
        for border in list(self.doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        normal=self.doc.styles['Normal'];normal.font.size=Pt(11)
        normal.paragraph_format.line_spacing=1.4;normal.paragraph_format.space_after=Pt(7)
        self.doc.styles['Title'].font.size=Pt(23)
        self.doc.styles['Title'].paragraph_format.space_after=Pt(10)
        for i,size in enumerate([18,15,12.5,11.5,11,11],1):
            style=self.doc.styles[f'Heading {i}'];style.font.size=Pt(size);style.font.bold=True
            style.paragraph_format.space_before=Pt(12 if i<3 else 8);style.paragraph_format.space_after=Pt(6)
            style.paragraph_format.keep_with_next=True
        self.doc.styles['Caption'].font.size=Pt(9);self.doc.styles['Caption'].font.italic=False
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        run=footer.add_run('墨记  ·  ');run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('777777')
        field=element('fldSimple',instr='PAGE');r=element('r');t=element('t');t.text='1';r.append(t);field.append(r);footer._p.append(field)
        self.numid=10

    def paragraph(self,parent,style=None,indent=0):
        p=parent.add_paragraph(style=style)
        p.paragraph_format.widow_control=True
        if indent:p.paragraph_format.left_indent=Pt(indent)
        return p

    def style_run(self,run,style):
        if style.get('bold'):run.bold=True
        if style.get('italic'):run.italic=True
        if style.get('strike'):run.font.strike=True
        if style.get('code'):
            run.font.name='Menlo';run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Heiti SC')
            run.font.size=Pt(9 if style.get('blockcode') else 10)
            if not style.get('blockcode'):shading(run._r.get_or_add_rPr(),'F0F1F3')
        if style.get('link'):run.font.color.rgb=RGBColor.from_string('42668E');run.underline=True

    def active(self,a,b):return [n for n in self.notes if n['start']<b and n['end']>a]
    def anchor(self,run,a,b):
        notes=self.active(a,b)
        for n in notes:self.anchors[n['id']].append(run)
        if notes:shading(run._r.get_or_add_rPr(),PALETTE[notes[-1]['color']])
        return notes

    def link(self,p,runs,url=None,anchor=None):
        if not runs:return
        if url and not re.match(r'^(?:https?://|codex:)',url):return
        link=element('hyperlink',history=1)
        if anchor:link.set(qn('w:anchor'),anchor)
        elif url:link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
        else:return
        first=runs[0]._r;first.addprevious(link)
        for run in runs:link.append(run._r)

    def marker(self,p,n):
        if n['id'] in self.marked:return
        self.marked.add(n['id']);r=p.add_run(f"[{self.numbers[n['id']]}]")
        r.font.superscript=True;r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
        self.link(p,[r],anchor='note_'+n['id'].replace('-',''))

    def finish_range(self,p,a,b):
        for n in self.notes:
            if a<n['end']<=b:self.marker(p,n)

    def text(self,p,pieces,style):
        for piece in pieces:
            a,b,value=piece['a'],piece['b'],piece['value']
            linear=b-a==len(value.encode('utf-16-le'))//2
            bounds=sorted({a,b,*[x for n in self.notes for x in [n['start'],n['end']] if a<x<b]}) if linear else [a,b]
            for x,y in zip(bounds,bounds[1:]):
                part=slice16(value,x-a,y-a) if linear else value
                # Host artifact references are readable file labels in documents.
                part=re.sub(r':codex-file-citation\{path="([^"]+)"[^}]*\}',lambda m:Path(m[1]).name,part)
                run=p.add_run(clean(part));self.style_run(run,style);self.anchor(run,x,y)
                if not self.quote_mode:self.finish_range(p,x,y)

    def image_run(self,p,node,data,width,math=False):
        if data.get('error') or not data.get('blob'):
            r=p.add_run('［'+data.get('error','图片暂时不可用')+'］');r.italic=True
            if node:
                a,b=span(node);self.anchor(r,a,b)
                if not self.quote_mode:self.finish_range(p,a,b)
            return
        image=base64.b64decode(data['blob'])
        if math:w,h=data['width'],data['height']
        else:w,h=data.get('width',600)*.75,data.get('height',400)*.75
        scale=min(1,width/max(w,1),(480 if not math else 400)/max(h,1));w*=scale;h*=scale
        notes=self.active(*span(node)) if node else []
        if notes:
            im=Image.open(BytesIO(image)).convert('RGBA');color='#'+PALETTE[notes[-1]['color']]
            if math:
                base=Image.new('RGBA',im.size,color);base.alpha_composite(im);im=base
            else:im=ImageOps.expand(im,border=8,fill=color)
            out=BytesIO();im.save(out,format='PNG');image=out.getvalue()
        r=p.add_run();shape=r.add_picture(BytesIO(image),width=Pt(max(w,1)),height=Pt(max(h,1)))
        shape._inline.docPr.set('descr',clean(data.get('latex') or data.get('filename') or '对话图片'))
        if math and data.get('depth') and node and node['type']=='inlineMath':
            r._r.get_or_add_rPr().append(element('position',val=-round(data['depth']*scale*2)))
        if node:
            a,b=span(node);self.anchor(r,a,b)
            if not self.quote_mode:self.finish_range(p,a,b)

    def image_data(self,node):
        ref=node if node['type']=='image' else self.definitions.get(node.get('identifier'),{})
        url=ref.get('url','');data=self.message.get('resources',{}).get(url,{'error':'图片文件暂时不可用'})
        return url,{**data,'filename':node.get('alt') or data.get('filename') or '图片'}

    def inline(self,p,node,style=None,width=None):
        style=style or {};width=width or self.width;kind=node['type']
        if kind in ('text','html','inlineCode','code'):
            self.text(p,node.get('map',[]),{**style,**({'code':True,'blockcode':kind=='code'} if kind in ('code','inlineCode') else {})});return
        if kind in ('math','inlineMath'):
            self.image_run(p,node,node['formula'],width,math=True);return
        if kind in ('image','imageReference'):
            _,data=self.image_data(node);self.image_run(p,node,data,width);return
        if kind=='break':p.add_run().add_break();return
        if kind=='footnoteReference':
            r=p.add_run('['+node.get('label',node['identifier'])+']');r.font.superscript=True
            a,b=span(node);self.anchor(r,a,b);return
        if kind=='strong':style={**style,'bold':True}
        if kind=='emphasis':style={**style,'italic':True}
        if kind=='delete':style={**style,'strike':True}
        if kind in ('link','linkReference'):
            definition=node if kind=='link' else self.definitions.get(node.get('identifier'),{})
            previous=len(p.runs)
            for child in node.get('children',[]):self.inline(p,child,{**style,'link':True},width)
            for run in p.runs[previous:]:self.link(p,[run],url=definition.get('url'))
            return
        for child in node.get('children',[]):self.inline(p,child,style,width)

    def numbering(self,ordered,start,indent):
        self.numid+=1;id_=self.numid
        root=self.doc.part.numbering_part.element
        abstract=element('abstractNum',abstractNumId=id_);abstract.append(element('multiLevelType',val='singleLevel'))
        level=element('lvl',ilvl=0);level.extend([element('start',val=start or 1),element('numFmt',val='decimal' if ordered else 'bullet'),element('lvlText',val='%1.' if ordered else '•'),element('lvlJc',val='left')])
        pp=element('pPr');pp.append(element('ind',left=round((indent+16)*20),hanging=240));level.append(pp);abstract.append(level);root.append(abstract)
        num=element('num',numId=id_);num.append(element('abstractNumId',val=id_));root.append(num);return id_

    def block(self,parent,node,width=None,indent=0):
        width=width or self.width;kind=node['type']
        if kind in ('definition','footnoteDefinition'):return
        if kind=='list':
            number=self.numbering(node.get('ordered'),node.get('start'),indent)
            for item in node['children']:
                first=True
                for child in item.get('children',[]):
                    if child['type']=='paragraph':
                        p=self.paragraph(parent,indent=indent+16)
                        if first:
                            if isinstance(item.get('checked'),bool):p.add_run('[x] ' if item['checked'] else '[ ] ')
                            else:
                                pp=p._p.get_or_add_pPr();np=element('numPr');np.extend([element('ilvl',val=0),element('numId',val=number)]);pp.append(np)
                        self.inline(p,child,width=width-16);first=False
                    else:self.block(parent,child,width-16,indent+16)
            return
        if kind=='table':self.table(parent,node,width,indent);return
        if kind=='blockquote':
            for child in node.get('children',[]):self.block(parent,child,width-14,indent+14)
            return
        if kind in ('paragraph','heading','code','math'):
            children=node.get('children',[])
            p=self.paragraph(parent,f"Heading {min(node.get('depth',2),6)}" if kind=='heading' else None,indent)
            if kind=='code':
                shading(p._p.get_or_add_pPr(),'F3F4F6');p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(10)
            single_image=kind=='paragraph' and len(children)==1 and children[0]['type'] in ('image','imageReference')
            if single_image:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if kind=='math':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(11)
            self.inline(p,node,width=width)
            if single_image:
                child=children[0];url,data=self.image_data(child)
                cap=self.paragraph(parent,'Caption',indent);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
                run=cap.add_run(clean(data['filename']));run.font.color.rgb=RGBColor.from_string('777777')
                if data.get('error') and url.startswith(('http://','https://')):self.link(cap,[run],url=url)
            return
        if kind=='thematicBreak':
            p=self.paragraph(parent);borders=element('pBdr');borders.append(element('bottom',val='single',sz=4,space=5,color='D9D9D9'));p._p.get_or_add_pPr().append(borders);return
        for child in node.get('children',[]):self.block(parent,child,width,indent)

    def table(self,parent,node,width,indent):
        rows=node.get('children',[])
        if not rows:return
        count=max(len(r['children']) for r in rows)
        weights=[4.0]*count
        for row in rows:
            for i,cell in enumerate(row['children']):
                def size(n):return len(n.get('value',''))+sum(size(c) for c in n.get('children',[]))
                weights[i]=max(weights[i],min(size(cell),42)**.65)
        sizes=[width*w/sum(weights) for w in weights]
        table=parent.add_table(rows=len(rows),cols=count);table.autofit=False
        pr=table._tbl.tblPr;pr.append(element('tblLayout',type='fixed'))
        borders=element('tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:borders.append(element(side,val='single',sz=5,color='D9D9D9'))
        pr.append(borders)
        margins=element('tblCellMar')
        for side,val in [('top',100),('bottom',100),('left',110),('right',110)]:margins.append(element(side,w=val,type='dxa'))
        pr.append(margins)
        for i,w in enumerate(sizes):table.columns[i].width=Pt(w)
        for i,row in enumerate(rows):
            if i==0:table.rows[i]._tr.get_or_add_trPr().append(element('tblHeader',val='true'))
            for j in range(count):
                cell=table.cell(i,j);cell.width=Pt(sizes[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.25
                if i==0:shading(cell._tc.get_or_add_tcPr(),self.header_fill)
                elif i%2==0:shading(cell._tc.get_or_add_tcPr(),'FAFAFB')
                align=node.get('align',[])
                if j<len(align) and align[j] in ('left','right','center'):p.alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'right':WD_ALIGN_PARAGRAPH.RIGHT,'center':WD_ALIGN_PARAGRAPH.CENTER}[align[j]]
                if j<len(row['children']):self.inline(p,row['children'][j],{'bold':i==0},sizes[j]-12)
        after=self.paragraph(parent);after.paragraph_format.space_after=Pt(2);after.paragraph_format.space_before=Pt(0);after.paragraph_format.line_spacing=Pt(2);after.add_run().font.size=Pt(2)

    def quote(self,n):
        message=next(m for m in self.item['messages'] if m['id']==n['message_id'])
        self.message=message;self.definitions={c['identifier']:c for c in message['ast']['children'] if c['type']=='definition'}
        p=self.paragraph(self.doc);p.paragraph_format.left_indent=Pt(10)
        run=p.add_run('原句：');run.font.color.rgb=RGBColor.from_string('777777')
        old_notes=self.notes;self.notes=[];self.quote_mode=True
        def walk(node):
            a,b=span(node)
            if a>=n['end'] or b<=n['start']:return
            kind=node['type']
            if kind in ('math','inlineMath'):self.image_run(p,node,node['formula'],self.width-45,math=True)
            elif kind in ('image','imageReference'):p.add_run('图片 '+node.get('alt',''))
            elif node.get('map'):
                for piece in node['map']:
                    x,y=max(n['start'],piece['a']),min(n['end'],piece['b'])
                    if x>=y:continue
                    linear=piece['b']-piece['a']==len(piece['value'].encode('utf-16-le'))//2
                    value=slice16(piece['value'],x-piece['a'],y-piece['a']) if linear else piece['value']
                    p.add_run(clean(value)).italic=True
            else:
                for child in node.get('children',[]):walk(child)
        walk(message['ast']);self.notes=old_notes;self.quote_mode=False

    def write(self,documents):
        if documents:
            self.doc.core_properties.title=clean(documents[0]['document']['title'])
            self.doc.core_properties.author='墨记';self.doc.core_properties.language='zh-CN'
        for di,item in enumerate(documents):
            self.item=item
            if di:self.doc.add_page_break()
            notes=sorted(item['annotations'],key=lambda n:(next(i for i,m in enumerate(item['messages']) if m['id']==n['message_id']),n['start'],n['end'],n['created_at']))
            self.numbers={n['id']:i+1 for i,n in enumerate(notes)}
            self.doc.add_paragraph(clean(item['document']['title']),'Title')
            intro=self.doc.add_paragraph(f"对话原文与批注 · {len(item['messages'])} 段原文 · {len(notes)} 处重点")
            intro.runs[0].font.size=Pt(9);intro.runs[0].font.color.rgb=RGBColor.from_string('777777')
            self.marked=set();self.anchors={n['id']:[] for n in notes};self.quote_mode=False
            for i,message in enumerate(item['messages']):
                self.message=message;self.notes=[n for n in notes if n['message_id']==message['id']]
                self.definitions={c['identifier']:c for c in message['ast']['children'] if c['type']=='definition'}
                heading=self.doc.add_paragraph(f"{ROLES.get(message['role'],'原文')}  {i+1:02d}",'Heading 3')
                self.block(self.doc,message['ast'])
                for attachment in message.get('attachments',[]):
                    p=self.paragraph(self.doc);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;self.image_run(p,None,attachment,self.width)
                    cap=self.doc.add_paragraph(clean(attachment.get('filename','对话图片')),'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for definition in [c for c in message['ast']['children'] if c['type']=='footnoteDefinition']:
                    p=self.paragraph(self.doc);p.add_run('['+definition['identifier']+'] ')
                    self.inline(p,definition,width=self.width)
                for n in self.notes:
                    selected=self.anchors[n['id']]
                    if not selected:
                        r=heading.add_run(f" [{self.numbers[n['id']]}]");selected=[r]
                    self.doc.add_comment([selected[0],selected[-1]],text=f"批注 {self.numbers[n['id']]}\n"+clean(n['comment'] or '仅高亮')+('\n标签：'+'、'.join(n['tags']) if n['tags'] else ''),author='墨记',initials='墨记')
                    if n['id'] not in self.marked:self.marker(heading,n)
            if notes:
                self.doc.add_paragraph('批注清单','Heading 1')
                for n in notes:
                    p=self.doc.add_paragraph(f"批注 {self.numbers[n['id']]}",'Heading 3')
                    self.bookmark+=1;p._p.append(element('bookmarkStart',id=self.bookmark,name='note_'+n['id'].replace('-','')));p._p.append(element('bookmarkEnd',id=self.bookmark))
                    self.quote(n);self.doc.add_paragraph(clean(n['comment'] or '仅高亮'))
                    if n['tags']:
                        p=self.doc.add_paragraph('标签：'+'、'.join(n['tags']));p.runs[0].font.size=Pt(9)
        out=BytesIO();self.doc.save(out);return out.getvalue()


def build(documents,appearance=None):
    documents=prepare(documents)
    warnings=[value['error'] for item in documents for message in item['messages']
        for value in [*message.get('resources',{}).values(),*message.get('attachments',[])] if value.get('error')]
    return WordWriter(appearance).write(documents),list(dict.fromkeys(warnings))


def docx_bytes(documents,appearance=None):return build(documents,appearance)[0]


def pdf_from_docx(data):
    config=json.loads((ROOT/'export-runtime.json').read_text());soffice=config.get('soffice')
    if not soffice or not Path(soffice).is_file():raise ValueError('PDF 排版环境未配置，请更新墨记')
    with tempfile.TemporaryDirectory(prefix='notes-export-',dir='/private/tmp' if os.path.isdir('/private/tmp') else None) as directory:
        folder=Path(directory);doc=folder/'document.docx';doc.write_bytes(data);profile=folder/'profile';profile.mkdir()
        env=os.environ.copy();env.update(TMPDIR=directory,TEMP=directory,TMP=directory,XDG_CONFIG_HOME=str(profile/'config'),XDG_CACHE_HOME=str(profile/'cache'))
        fonts=Path(config['python']).parents[2]/'native/poppler/poppler/etc/fonts'
        if (fonts/'fonts.conf').is_file():env.update(FONTCONFIG_FILE=str(fonts/'fonts.conf'),FONTCONFIG_PATH=str(fonts))
        options={'ExportNotes':{'type':'boolean','value':'true'},'ExportNotesInMargin':{'type':'boolean','value':'false'},'UseTaggedPDF':{'type':'boolean','value':'true'},'ExportBookmarks':{'type':'boolean','value':'true'}}
        command=[soffice,'-env:UserInstallation='+profile.as_uri(),'--headless','--norestore','--convert-to','pdf:writer_pdf_Export:'+json.dumps(options),'--outdir',directory,str(doc)]
        result=subprocess.run(command,capture_output=True,text=True,env=env,timeout=75)
        pdf=folder/'document.pdf'
        if not pdf.is_file() or pdf.stat().st_size<100:raise ValueError('PDF 排版失败，请稍后重试。'+result.stderr[-180:])
        from pypdf import PdfReader,PdfWriter
        from pypdf.generic import NameObject,ArrayObject,FloatObject,BooleanObject
        reader=PdfReader(pdf);writer=PdfWriter();writer.clone_document_from_reader(reader)
        for page in writer.pages:
            used=[];right=float(page.mediabox.right);top=float(page.mediabox.top)
            for ref in page.get('/Annots',[]):
                annotation=ref.get_object()
                if annotation.get('/Subtype')=='/Popup':annotation[NameObject('/Open')]=BooleanObject(False)
                if annotation.get('/Subtype')!='/Text':continue
                rect=annotation.get('/Rect',[0,0,0,0]);y=max(32,min(top-48,float(rect[1])))
                for _ in range(max(1,int((top-70)/20))):
                    if not any(abs(y-p)<20 for p in used):break
                    y=y-20 if y>52 else top-48
                used.append(y)
                annotation[NameObject('/Rect')]=ArrayObject([FloatObject(x) for x in [right-34,y,right-18,y+16]])
                annotation[NameObject('/Name')]=NameObject('/Comment');annotation[NameObject('/Open')]=BooleanObject(False)
        stream=BytesIO();writer.write(stream);return stream.getvalue()
